"""Automatic, source-addressed imports. No model fitting and no network access.

The offline adapter recognizes common economy columns. Optional LLM facts use
this same intermediate contract, then the original canonical schema/vectorizer.
Unknown cells and unsupported semantics remain in the original files.
"""
from __future__ import annotations
import base64, csv, io, json, re, hashlib, math, copy, sys, struct
from pathlib import Path
from datetime import date, datetime
from functools import lru_cache

FIELDS = {
 'name': ['name','name_column','item','items','item_name','resource','resource_name','title','sword_name','satchel_name','entity_name'],
 'kind': ['entity_type','entity','type','kind'],
 'category': ['category','resource_category','item_type'],
 'currency': ['currency','currency_name','price_currency','payment_resource'],
 'buy_price': ['buy_price','purchase_price','purchase_cost','buy','price_to_buy','cost_to_buy'],
 'sell_price': ['sell_price','sell','sale_price','resale_value','price_to_sell','cost_to_sell'],
 'merchant_price': ['sells_for','sells_at','asking_price'],
 'cost_amount': ['cost','price','cost_amount','upgrade_cost','amount'],
 'source_rate': ['source_rate','earning_rate','production_rate','income_rate','generation_rate'],
 'reward_amount': ['reward','reward_amount','output_amount','yield'],
 'timer_duration': ['cooldown','duration','timer_duration','time','crafting_time'],
 'capacity': ['capacity','weight_capacity','carrying_capacity','max_weight','storage_capacity','stack_size','max_stack','stack'],
 'weight': ['weight','item_weight','mass'],
 'drop_probability': ['probability','drop_chance','drop_probability','chance'],
 'conversion_ratio': ['conversion_ratio','exchange_ratio'],
 'frequency': ['frequency'], 'visibility':['visibility'], 'recurrence':['recurrence'],
 'dismantle': ['dismantle_components','dismantling','dismantle','salvage','salvage_outputs'],
 'ingredients': ['ingredients','recipe','crafting_requirements','components','inputs'],
 'upgrade_material': ['upgrade','upgrade_material','upgrade_materials','upgrade_requirement','upgrade_requirements','tempering_material'],
 'merchant': ['merchant','trader','vendor','shop','seller'],
 'inventory': ['inventory','stored_in','storage','container'],
}
NUMERIC = ['buy_price','sell_price','merchant_price','cost_amount','source_rate','reward_amount','timer_duration','capacity','weight','drop_probability','conversion_ratio','frequency','visibility','recurrence']
KINDS = ['resource','source','sink','converter','inventory','trader']
PARAMETERS = ['cost_amount','source_rate','reward_amount','timer_duration','capacity','drop_probability','conversion_ratio','frequency','visibility','recurrence']
RESOURCE_PROPERTIES = ['scarcity_level','convertibility','tradability','perishability','progression_relevance','power_relevance','cosmetic_relevance','monetization_relevance','time_relevance','status_relevance','competitive_relevance']
JSON_COLUMNS = {'all_infobox_fields_json','extra_fields_json'}

def nested_cells(table,ri):
 """Expose explicit infobox scalars with reversible JSON locations, never prose guesses."""
 result=[]
 for ci,h in enumerate(table['headers']):
  if norm(h) not in JSON_COLUMNS:continue
  try:data=json.loads(table['values'][ri][ci] or '{}')
  except (ValueError,TypeError):continue
  if not isinstance(data,dict):continue
  row,col=table['coordinates'][ri][ci]
  for key,value in data.items():
   if heading(key) not in NUMERIC or isinstance(value,(dict,list)):continue
   pointer=str(key).replace('~','~0').replace('/','~1')
   result.append({'source_id':f'{table["id"]}!R{row}C{col}#/{pointer}','file_index':table['file_index'],'sheet':table['sheet'],'row':row,'column':col,'header':h,'original':value,'value':value,'table_id':table['id'],'record_index':ri,'file_name':table['file_name'],'json_path':[key],'nested_header':key})
 return result

def norm(s):return _norm(str(s))
@lru_cache(maxsize=4096)
def _norm(s):return re.sub(r'[^\w]+','_',s.casefold()).strip('_')
def ident(prefix,s): return prefix+'_'+hashlib.sha256(str(s).encode()).hexdigest()[:16]
def simple(v):
 if isinstance(v,(date,datetime)): return v.isoformat()
 if isinstance(v,float) and not math.isfinite(v):return None
 return v

def decode_csv(blob):
 if blob.startswith((b'\xff\xfe',b'\xfe\xff')):return blob.decode('utf-16'),'utf-16'
 try:return blob.decode('utf-8-sig'),'utf-8-sig' if blob.startswith(b'\xef\xbb\xbf') else 'utf-8'
 except UnicodeDecodeError:return blob.decode('cp1252'),'cp1252'

def csv_dialect(text,ext):
 try:return csv.Sniffer().sniff(text[:65536],delimiters=',;\t')
 except csv.Error:
  first=next((line for line in text.splitlines() if line.strip()),'')
  class D(csv.excel):pass
  D.delimiter=max([',',';','\t'],key=first.count) if any(c in first for c in ',;\t') else ('\t' if ext=='.tsv' else ',')
  return D

def heading(s):return _heading(str(s))

def resolved_field(table,column):
 # Recognized source meanings take precedence over mappings saved by older apps.
 return heading(table['headers'][column]) or table.get('column_mapping',{}).get(str(column))

@lru_cache(maxsize=4096)
def _heading(s):
 n=norm(s)
 if n in RESOURCE_PROPERTIES:return n
 if n in ['functional_roles','roles']:return 'functional_roles'
 for field,aliases in FIELDS.items():
  if n in aliases:return field
 # Header units are optional and independent from column names.
 n=norm(re.sub(r'\([^)]*\)|\[[^]]*\]','',str(s)))
 for field,aliases in FIELDS.items():
  if n in aliases:return field
 for field,aliases in FIELDS.items():
  if field=='name':continue
  if any(n.startswith(a+'_') and re.search(r'coin|gold|gem|usd|second|minute|hour|percent|kg|unit',n[len(a):]) for a in aliases):return field
 return None

def amount(value):
 """Recognize one scalar with optional units. Never coerce ranges to numbers."""
 if value is None or isinstance(value,bool):return None
 if isinstance(value,(float,int)):return float(value) if math.isfinite(value) and value>=0 else None
 s=str(value).strip()
 if s.startswith('=') or not s:return None
 s=re.sub(r'(?<=\d),(?=\d{3}(?:\D|$))','',s)
 matches=list(re.finditer(r'(?<![\w.-])\d+(?:\.\d+)?(?:[eE][+-]?\d+)?|(?<![\w.-])\.\d+',s))
 if len(matches)!=1 or re.search(r'\d\s*[-–—/]\s*\d',s):return None
 m=matches[0]
 rest=s[:m.start()]+s[m.end():]
 if re.search(r'\d|[=+<>]|^-',rest):return None
 value=float(m[0]);return value if math.isfinite(value) else None

def unit_for(header,raw,field,currency=None):
 text=f'{header} {raw or ""}'.casefold()
 if field in ['drop_probability','visibility']:return 'percent' if '%' in text or 'percent' in text else '0–1'
 if field in ['buy_price','sell_price','merchant_price','cost_amount']:
  for pattern,name in [(r'\bgems?\b','Gems'),(r'\bcoins?\b','Coins'),(r'\bgold\b','Gold'),(r'\bbells\b','Bells'),(r'\bcredits?\b','Credits'),(r'\btokens?\b','Tokens')]:
   if re.search(pattern,text):return name
  return str(currency) if currency else 'Unspecified currency'
 bracket=re.search(r'\(([^)]+)\)|\[([^]]+)\]',str(header))
 if bracket:return bracket[1] or bracket[2]
 if field=='weight' or (field=='capacity' and 'weight' in text):return 'kg' if re.search(r'\bkg\b',text) else 'weight units'
 if field=='capacity':return 'units'
 if field=='timer_duration':
  for u in ['seconds','minutes','hours','days']:
   if u[:-1] in text:return u
  return 'time units (unspecified)'
 if field=='source_rate':
  for u in ['second','minute','hour','day']:
   if u in text:return f'{currency or "resource"} per {u}'
  return f'{currency or "resource"} per source time unit'
 return 'per action' if field in ['reward_amount','frequency'] else 'ratio' if field=='conversion_ratio' else 'source units'

def parse_files(files):
 # csv's default 128 KiB field limit is an implementation default, not an input policy.
 limit=sys.maxsize
 while True:
  try:csv.field_size_limit(limit);break
  except OverflowError:limit//=10
 tables=[];canonical=[]
 for fi,f in enumerate(files):
  blob=base64.b64decode(f['data']);ext=Path(f['name']).suffix.lower();sheets=[];record_paths={}
  if ext in ['.csv','.tsv']:
   text,encoding=decode_csv(blob);dialect=csv_dialect(text,ext)
   sheets=[('Data',list(csv.reader(io.StringIO(text),dialect)),None)]
  elif ext in ['.xlsx','.xlsm']:
   import openpyxl
   wb=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=False)
   cached=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
   for ws in wb.worksheets:
    sheets.append((ws.title,[[simple(v) for v in r] for r in ws.iter_rows(values_only=True)],[[simple(v) for v in r] for r in cached[ws.title].iter_rows(values_only=True)]))
   wb.close();cached.close()
  elif ext=='.xls':
   import xlrd
   wb=xlrd.open_workbook(file_contents=blob,on_demand=True)
   for wi,ws in enumerate(wb.sheets()):
    values=[ws.row_values(i) for i in range(ws.nrows)];original=copy.deepcopy(values)
    # xlrd exposes cached results, not formula text. Mark FORMULA records so
    # their results cannot be mistaken for directly editable source constants.
    pos=wb._sh_abs_posn[wi];mem=wb.mem
    while pos+4<=len(mem):
     opcode,length=struct.unpack_from('<HH',mem,pos);pos+=4
     if opcode==0x000A:break
     if opcode in [0x0006,0x0206,0x0406] and length>=6:
      r,c=struct.unpack_from('<HH',mem,pos)
      if r<len(original) and c<len(original[r]):original[r][c]='= [formula retained in original XLS; showing cached result]'
     pos+=length
    sheets.append((ws.name,original,values))
   wb.release_resources()
  elif ext=='.json':
   obj=json.loads(blob)
   if isinstance(obj,dict) and 'entities' in obj and 'resources' in obj:
    canonical.append({'file_index':fi,'economy':obj});continue
   records=obj if isinstance(obj,list) else obj.get('items',obj.get('data')) if isinstance(obj,dict) else None
   if not isinstance(records,list) or not all(isinstance(r,dict) for r in records):
    raise ValueError(f'{f["name"]}: use a table of records or canonical economy JSON.')
   headers=list(dict.fromkeys(k for r in records for k in r))
   sheets=[('Data',[headers]+[[r.get(k) for k in headers] for r in records],None)]
  else:raise ValueError(f'{f["name"]}: use CSV, TSV, Excel (.xlsx, .xlsm, .xls), or JSON.')
  for sheet,matrix,cached in sheets:
   nonempty=[i for i,row in enumerate(matrix) if any(v is not None and str(v).strip() for v in row)]
   if not nonempty:continue
   # Recognize key/value documents as well as conventional column tables.
   vertical=len(nonempty)>1 and sum(heading(matrix[i][0]) is not None for i in nonempty if matrix[i])>len(nonempty)/2 and max(len([v for v in matrix[i] if v is not None and str(v).strip()]) for i in nonempty)<=2 and heading(matrix[nonempty[0]][1] if len(matrix[nonempty[0]])>1 else '') is None
   if vertical:
    headers=[str(matrix[i][0]) for i in nonempty];rows=[[matrix[i][1] if len(matrix[i])>1 else None for i in nonempty]]
    coords=[[(i+1,2) for i in nonempty]];header_row=nonempty[0]+1
    values=[[cached[i][1] if cached and len(cached[i])>1 else v for i,v in zip(nonempty,rows[0])]]
   else:
    # Prefer the first recognizable header row; preserve all later rows/columns.
    hi=next((i for i in nonempty if sum(heading(v) is not None for v in matrix[i])>=2),nonempty[0])
    width=max(len(row) for row in matrix[hi:]);base=list(matrix[hi])+[None]*(width-len(matrix[hi]));counts={};headers=[]
    for ci,h in enumerate(base):
     h=str(h) if h is not None and str(h).strip() else f'Column {ci+1}';counts[h]=counts.get(h,0)+1
     headers.append(h if counts[h]==1 else f'{h} [{counts[h]}]')
    indices=[i for i in nonempty if i>hi];rows=[list(matrix[i])+[None]*(width-len(matrix[i])) for i in indices]
    coords=[[(i+1,c+1) for c in range(width)] for i in indices];header_row=hi+1
    values=[[(cached[i][c] if cached and i<len(cached) and c<len(cached[i]) else None) if isinstance(v,str) and v.startswith('=') else v for c,v in enumerate(row)] for i,row in zip(indices,rows)]
   keys={field:next((h for h in headers if heading(h)==field),'') for field in FIELDS}
   mapping={k:keys.get(k,'') for k in ['name','entity_type','resource','cost_amount','source_rate','capacity']}
   tables.append({'id':f'{fi}:{sheet}','file_index':fi,'file_name':f['name'],'sheet':sheet,'headers':headers,'rows':rows,'values':values,'coordinates':coords,'header_row':header_row,'preamble':matrix[:header_row-1],'vertical':vertical,'mapping':mapping,'units':{},'kind':'resource','include':True})
 if not tables and not canonical:raise ValueError('No populated sheets or economy records were found.')
 return {'tables':tables,'canonical':canonical[0] if len(canonical)==1 and not tables else None,'canonical_files':canonical}

def cells(parsed):
 result={}
 for t in parsed['tables']:
  for ri,row in enumerate(t['rows']):
   for ci,raw in enumerate(row):
    r,c=t['coordinates'][ri][ci];sid=f'{t["id"]}!R{r}C{c}'
    result[sid]={'source_id':sid,'file_index':t['file_index'],'sheet':t['sheet'],'row':r,'column':c,'header':t['headers'][ci],'original':raw,'value':t['values'][ri][ci],'table_id':t['id'],'record_index':ri,'file_name':t['file_name']}
   for cell in nested_cells(t,ri):result[cell['source_id']]=cell
 return result

def parts(value):
 result=[]
 for s in re.split(r'\s*\+\s*|;|\n',str(value or '')):
  m=re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*[x×*]?\s+(.+?)\s*',s) or re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*[x×*]\s*(.+?)\s*',s)
  if m:result.append({'resource':m[2].strip(),'quantity':float(m[1])})
  elif s.strip():result.append({'resource':s.strip(),'quantity':None})
 return result

def offline_facts(parsed):
 facts=[];notes=[]
 for t in parsed['tables']:
  fields={}
  for ci,h in enumerate(t['headers']):
   k=resolved_field(t,ci)
   if k and k not in fields:fields[k]=ci
  for ri,row in enumerate(t['rows']):
   vals=t['values'][ri]
   nested=nested_cells(t,ri);fallback={}
   for cell in nested:fallback.setdefault(heading(cell['nested_header']),cell)
   def v(k):
    value=vals[fields[k]] if k in fields else None
    return fallback[k]['value'] if value in [None,''] and k in fallback else value
   def sid(ci):
    r,c=t['coordinates'][ri][ci];return f'{t["id"]}!R{r}C{c}'
   rawname=v('name')
   if rawname is None:
    rawname=next((x for x in vals if isinstance(x,str) and x.strip() and not x.startswith('=') and amount(x) is None),None)
   if rawname is None:rawname=f'{t["sheet"]} row {t["coordinates"][ri][0][0]}'
   name=str(rawname).strip()
   if not name or (heading(name)=='name' and sum(heading(x) is not None for x in row)>=2):continue
   context=f'{t["sheet"]} {t["file_name"]} {v("category") or ""}'.casefold()
   kind=norm(v('kind') or '')
   if kind not in KINDS:
    kind='inventory' if v('capacity') not in [None,''] and any(s in context for s in ['satchel','bag','inventory','storage','container','wallet']) else 'source' if v('source_rate') not in [None,''] or v('reward_amount') not in [None,''] else 'resource'
   cat=[]
   for words,c in [(['sword','weapon','bow','dagger'],'weapon'),(['armor','armour','shield'],'armor'),(['material','wood','steel','ore'],'crafting_item'),(['cosmetic','skin'],'cosmetic_item'),(['currency','coin','gold'],'soft_currency')]:
    if any(w in context for w in words):cat.append(c)
   if v('category'):cat.append(norm(v('category')))
   attrs=[]
   for field in NUMERIC:
    if v(field) in [None,'']:continue
    use_nested=field in fallback and (field not in fields or vals[fields[field]] in [None,''])
    if use_nested:
     cell=fallback[field];raw=cell['value'];header=cell['nested_header'];source_id=cell['source_id']
    else:
     ci=fields[field];raw=vals[ci];header=t['headers'][ci];source_id=sid(ci)
    value=amount(raw);unit=unit_for(header,raw,field,v('currency'))
    if value is None:notes.append(f'{t["file_name"]} / {t["sheet"]}, {source_id}: {header} retained as text or unknown.')
    attrs.append({'field':field,'value':value,'unit':unit,'source_id':source_id})
   merchant=str(v('merchant') or '').strip()
   if not merchant and re.search(r'merchant|vendor|shop',t['sheet'],re.I):merchant=t['sheet']
   if not merchant and re.search(r'merchant|vendor|shop',Path(t['file_name']).stem,re.I):merchant=Path(t['file_name']).stem
   recipes=[]
   for key in ['dismantle','ingredients']:
    if v(key):recipes.append({'action':'dismantle' if key=='dismantle' else 'craft','inputs':[{'resource':name,'quantity':1.}] if key=='dismantle' else parts(v(key)),'outputs':parts(v(key)) if key=='dismantle' else [{'resource':name,'quantity':1.}],'source_ids':[sid(fields[key])]})
   requirements=[]
   if v('upgrade_material') and norm(v('upgrade_material')) not in {'none','na','n_a','not_applicable','cannot_be_upgraded'}:
    requirements=[{'materials':parts(v('upgrade_material')),'source_ids':[sid(fields['upgrade_material'])]}]
   properties=[]
   for key in RESOURCE_PROPERTIES:
    if key not in fields:continue
    raw=v(key);val=amount(raw)
    if str(raw).casefold() in ['true','yes']:val=1.
    elif str(raw).casefold() in ['false','no']:val=0.
    if isinstance(raw,str) and '%' in raw and val is not None:val/=100
    if val is not None and not 0<=val<=1:val=None
    properties.append({'name':key,'value':val,'source_ids':[sid(fields[key])],'reason':'Explicit source property; unknown when its scale is not stated.'})
   roles=[x.strip() for x in re.split(r'[,;]',str(v('functional_roles') or '')) if x.strip()]
   facts.append({'record_id':f'{t["id"]}:record:{ri}','name':name,'kind':kind,'categories':list(dict.fromkeys(cat)),'resource':str(v('currency') or name),'merchant':merchant,'inventory':str(v('inventory') or ''),'attributes':attrs,'properties':properties,'functional_roles':roles,'connections':[],'recipes':recipes,'source_ids':[sid(i) for i,x in enumerate(row) if x not in [None,'']],'notes':[]})
   facts[-1]['upgrade_requirements']=requirements
 return {'facts':facts,'notes':notes,'method':'offline_rules_v2'}

def validate_facts(document,parsed,source=None):
 """LLM output is data, never executable code or a ready-made model vector."""
 source=cells(parsed) if source is None else source;seen=set()
 if not isinstance(document,dict) or not isinstance(document.get('facts'),list):raise ValueError('Interpreter returned no fact records.')
 for f in document['facts']:
  if f['record_id'] in seen:raise ValueError('Interpreter returned duplicate record identifiers.')
  seen.add(f['record_id'])
  if f['kind'] not in KINDS or not f['name'].strip():raise ValueError('Invalid economy fact kind or name.')
  if not f['source_ids'] or any(s not in source for s in f['source_ids']):raise ValueError('Interpreter cited a missing source cell.')
  for a in f['attributes']:
   if a['field'] not in NUMERIC or a['source_id'] not in source:raise ValueError('Unsupported field or source cell in interpreted data.')
   cell=source[a['source_id']]
   if heading(cell.get('nested_header',cell.get('header','')))=='upgrade_material':raise ValueError('Upgrade-material quantities are not editable price or economy-rate parameters.')
   val=a['value'];raw=source[a['source_id']]['value']
   if val is not None:
    if isinstance(val,bool) or not isinstance(val,(int,float)) or not math.isfinite(val) or val<0:raise ValueError('Invalid numeric fact.')
    observed=amount(raw)
    if observed is None or not math.isclose(val,observed,rel_tol=1e-9,abs_tol=1e-9):raise ValueError('A numeric interpretation does not match its source cell. Offline interpretation retained.')
  for recipe in f['recipes']:
   if not recipe['source_ids'] or any(s not in source for s in recipe['source_ids']):raise ValueError('Recipe cites a missing source cell.')
   evidence_fields={heading(source[s].get('nested_header',source[s].get('header',''))) for s in recipe['source_ids']}
   if 'upgrade_material' in evidence_fields and evidence_fields<={'upgrade_material','name','kind','category'}:raise ValueError('An upgrade-material requirement does not establish a crafting or conversion recipe.')
   evidence=' '.join(str(source[s]['value']) for s in recipe['source_ids'])
   for side in ['inputs','outputs']:
    for item in recipe[side]:
     q=item['quantity']
     if q is not None and (isinstance(q,bool) or not isinstance(q,(int,float)) or not math.isfinite(q) or q<0):raise ValueError('Invalid recipe quantity.')
     if q not in [None,1,1.] and not any(math.isclose(q,float(x)) for x in re.findall(r'\d+(?:\.\d+)?',evidence)):raise ValueError('Recipe quantity has no source evidence.')
  for prop in f.get('properties',[]):
   if prop['name'] not in RESOURCE_PROPERTIES or not prop['source_ids'] or any(s not in source for s in prop['source_ids']):raise ValueError('Resource property has an unknown name or source.')
   v=prop['value']
   if v is not None and (isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v) or not 0<=v<=1):raise ValueError('Resource properties must be in [0,1] or unknown.')
  for flow in f.get('connections',[]):
   if flow['direction'] not in ['input','output'] or not flow['source_ids'] or any(s not in source for s in flow['source_ids']):raise ValueError('Connection has an invalid direction or source.')
   for k in ['flow_rate','probability','delay']:
    v=flow[k]
    if v is not None and (isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v) or v<0 or (k=='probability' and v>1)):raise ValueError('Connection quantities are invalid.')
    if v is not None and not any(amount(source[s]['value']) is not None and math.isclose(v,amount(source[s]['value'])) for s in flow['source_ids']):raise ValueError('Connection quantity needs direct scalar evidence.')
 return document

def build_economy(parsed,project_id,title,document=None):
 document=document or offline_facts(parsed);source=cells(parsed);resources={};entities=[];links=[];relations=[];bindings=[];notes=list(document.get('notes',[]));seenlinks=set();wallets={};weights={};bags=[];roundtrips={};row_attributes={};constraints={};constraint_conflicts=[];fact_entities={};property_conflicts=set();seenrecipes={};seenoffers={}
 def currency_categories(name):
  return ['soft_currency'] if norm(name) in ['coin','coins','gold','bells','credits'] else ['currency']
 def resource(name,categories=None,provenance=None):
  rid=ident('r',norm(name));cats=categories or []
  if rid not in resources:resources[rid]={'resource_id':rid,'name':name,'categories':list(cats),'properties':{},'provenance':{'source_ids':provenance or [],'mapping_confidence':.7}}
  else:
   resources[rid]['categories']=list(dict.fromkeys(resources[rid]['categories']+cats));resources[rid]['provenance']['source_ids']=list(dict.fromkeys(resources[rid]['provenance']['source_ids']+(provenance or [])))
  return rid
 def entity(eid,name,kind,rids,params=None,roles=None,src=None):
  obj={'entity_id':eid,'name':name,'entity_type':kind,'affected_resources':list(dict.fromkeys(rids)),'functional_roles':roles or [],'parameters':params or {},'constraints':{'can_modify':False,'parameter_units':{}},'provenance':{'source_ids':src or [],'mapping_confidence':.7,'notes':'Automatic import; inspect the source evidence and assumptions.'}}
  entities.append(obj);return obj
 def holding(rid):
  if rid not in wallets:wallets[rid]=entity('hold_'+rid,'Hold '+resources[rid]['name'],'inventory',[rid],roles=['resource_holding'])
  return wallets[rid]
 def link(a,b,rid):
  key=(a['entity_id'],b['entity_id'],rid)
  if key in seenlinks:return
  seenlinks.add(key);links.append({'link_id':ident('l',repr(key)),'source_entity_id':key[0],'target_entity_id':key[1],'resource_id':rid,'flow_rate':None,'probability':None,'delay':None})
 def bind(ent,param,a):
  val=a['value'];unit=a.get('unit') or 'source units';scale=100 if unit=='percent' else 1
  if val is None:return
  val=val/scale
  if param in ['drop_probability','visibility'] and not 0<=val<=1:
   notes.append(f'{ent["name"]}: {param} remains unknown because its scale is ambiguous.');return
  ent['parameters'][param]=val;ent['constraints']['parameter_units'][param]='0–1' if scale==100 else unit
  ent['constraints'][param+'_min']=val;ent['constraints'][param+'_max']=val
  cell=source.get(a['source_id']);direct=cell is not None and not (isinstance(cell['original'],str) and cell['original'].startswith('='))
  if direct:
   bindings.append({**cell,'entity_id':ent['entity_id'],'parameter':param,'scale':scale,'transform':'scalar','unit':unit})
   if cell.get('json_path'):
    # Duplicate infobox representations must receive the same accepted edit.
    for other in nested_by_row.get((cell['table_id'],cell['record_index']),[]):
     if other['source_id']!=cell['source_id'] and norm(other['nested_header'])==norm(cell['nested_header']) and other['value']==cell['value']:
      bindings.append({**other,'entity_id':ent['entity_id'],'parameter':param,'scale':scale,'transform':'scalar','unit':unit})
   if param in PARAMETERS and param!='reward_amount':
    ent['constraints']['can_modify']=True;ent['constraints'][param+'_min']=max(0,val*.8);ent['constraints'][param+'_max']=min(1,val*1.2) if param in ['visibility','drop_probability'] else val*1.2
 nested_by_row={}
 for cell in source.values():
  if cell.get('json_path'):nested_by_row.setdefault((cell['table_id'],cell['record_index']),[]).append(cell)
 for fact in document['facts']:
  name=fact['name'];attrs={a['field']:a for a in fact['attributes']};
  priced=[attrs[k] for k in ['buy_price','sell_price','merchant_price'] if k in attrs]
  if 'cost_amount' in attrs and any(a['source_id']==attrs['cost_amount']['source_id'] for a in priced):
   attrs.pop('cost_amount');notes.append('A redundant cost alias for the same price cell was removed.')
  ridsrc=fact['source_ids'];resource_name=fact['resource'] if fact['kind']!='resource' and fact['resource'] else name
  rid=resource(resource_name,fact['categories'] if resource_name==name else currency_categories(resource_name),ridsrc);ridholder=holding(rid);row_attributes[fact['record_id']]=attrs
  fact_entities[fact['record_id']]=ridholder
  for prop in fact.get('properties',[]):
   key=prop['name'];value=prop['value'];known=resources[rid]['properties'].get(key)
   if (rid,key) in property_conflicts:continue
   if value is not None and known is not None and value!=known:
    property_conflicts.add((rid,key));resources[rid]['properties'][key]=None;notes.append(f'{resource_name}: conflicting {key} values retained in evidence; canonical property is unknown.')
   elif value is not None:resources[rid]['properties'][key]=value
  if 'weight' in attrs and attrs['weight']['value'] is not None:weights.setdefault(rid,[]).append({**attrs['weight'],'inventory':fact.get('inventory','')})
  if fact['kind']=='inventory':
   ent=entity(ident('e',fact['record_id']+':inventory'),name,'inventory',[rid],roles=['inventory_capacity','storage'],src=ridsrc)
   fact_entities[fact['record_id']]=ent
   if 'capacity' in attrs:bind(ent,'capacity',attrs['capacity']);bags.append((ent,attrs['capacity']))
  elif fact['kind'] in ['source','sink','converter','trader'] and not ((priced or fact['recipes']) and not any(k in PARAMETERS for k in attrs)):
   target=rid
   ent=entity(ident('e',fact['record_id']+':entity'),name,fact['kind'],[target],roles=[fact['kind']],src=ridsrc)
   fact_entities[fact['record_id']]=ent
   for k,a in attrs.items():
    if k in PARAMETERS:bind(ent,k,a)
   if fact['kind']=='source':link(ent,holding(target),target)
   else:link(holding(target),ent,target)
  elif 'capacity' in attrs:
   bind(ridholder,'capacity',attrs['capacity']);ridholder['functional_roles'].append('storage')
  for field in ['buy_price','sell_price','merchant_price']+(['cost_amount'] if fact['kind']=='resource' else []):
   a=attrs.get(field)
   if not a or a['value'] is None:continue
   offer_key=(rid,a['source_id'],'sell' if field=='sell_price' and not fact['merchant'] else 'buy',norm(fact['merchant']))
   if offer_key in seenoffers:
    fact_entities[fact['record_id']]=seenoffers[offer_key];notes.append('A repeated offer from the same source cell was counted once.');continue
   currency=a['unit'] or 'Unspecified currency'
   if currency=='Unspecified currency':notes.append(f'{name}: price currency was not stated; kept separate as “Unspecified currency”.')
   crid=resource(currency,currency_categories(currency),[a['source_id']]);holder=holding(crid)
   sale=field=='sell_price' and not fact['merchant'];merchant=fact['merchant']
   # Match prepare.py: buy_price AND sell_price are trader.cost_amount.
   # Full conversion direction and merchant scope remain in relationship metadata.
   trade=entity(ident('e',fact['record_id']+':'+field),f'{merchant+" · " if merchant else ""}{"Sell" if sale else "Buy"} {name}','trader',[rid,crid],roles=['selling' if sale else 'buying'],src=[a['source_id']])
   seenoffers[offer_key]=trade
   trade['constraints']['source_field']=field
   bind(trade,'cost_amount',a)
   if sale:link(ridholder,trade,rid);link(trade,holder,crid)
   else:link(holder,trade,crid);link(trade,ridholder,rid)
   relations.append({'type':'exchange','entity_id':trade['entity_id'],'name':trade['name'],'inputs':[{'resource_id':rid if sale else crid,'quantity':1 if sale else a['value']}],'outputs':[{'resource_id':crid if sale else rid,'quantity':a['value'] if sale else 1}],'source_ids':[a['source_id']],'merchant':merchant or None})
   if not merchant:roundtrips.setdefault((rid,crid),{})['sell' if sale else 'buy']=(a,trade)
  for idx,recipe in enumerate(fact['recipes']):
   signature=(tuple(sorted(recipe['source_ids'])),tuple(sorted((norm(x['resource']),x['quantity']) for x in recipe['inputs'])),tuple(sorted((norm(x['resource']),x['quantity']) for x in recipe['outputs'])))
   if signature in seenrecipes:
    if fact['kind']=='converter':fact_entities[fact['record_id']]=seenrecipes[signature]
    notes.append('A repeated recipe from the same source cells was counted once.');continue
   inputs=[{'resource_id':resource(x['resource']), 'quantity':x['quantity']} for x in recipe['inputs']];outputs=[{'resource_id':resource(x['resource']), 'quantity':x['quantity']} for x in recipe['outputs']]
   ent=entity(ident('e',fact['record_id']+':recipe:'+str(idx)),f'{recipe["action"].capitalize()} {name}','converter',[x['resource_id'] for x in inputs+outputs],roles=['crafting' if recipe['action']=='craft' else 'dismantling'],src=recipe['source_ids'])
   seenrecipes[signature]=ent
   if fact['kind']=='converter':fact_entities[fact['record_id']]=ent
   for x in inputs:link(holding(x['resource_id']),ent,x['resource_id'])
   for x in outputs:link(ent,holding(x['resource_id']),x['resource_id'])
   relations.append({'type':'recipe','entity_id':ent['entity_id'],'name':ent['name'],'inputs':inputs,'outputs':outputs,'source_ids':recipe['source_ids']})
  for requirement in fact.get('upgrade_requirements',[]):
   materials=[{'resource_id':resource(x['resource'],provenance=requirement['source_ids']),'quantity':x['quantity']} for x in requirement['materials']]
   for material in materials:holding(material['resource_id'])
   if materials:relations.append({'type':'upgrade_requirement','name':f'Upgrade material for {name}','resource_id':rid,'materials':materials,'source_ids':requirement['source_ids'],'note':'Material requirement only. Missing quantities and the upgrade result remain unknown; no crafting flow is inferred.'})
  for k in ['drop_probability','timer_duration','frequency','visibility','recurrence']:
   if fact['kind']=='resource' and k in attrs:bind(ridholder,k,attrs[k])
  target=fact_entities[fact['record_id']];target['functional_roles']=list(dict.fromkeys(target['functional_roles']+fact.get('functional_roles',[])))
 # Optional LLM connections carry only explicit source-supported topology/quantities.
 named={}
 for fact in document['facts']:named.setdefault(norm(fact['name']),[]).append(fact_entities[fact['record_id']])
 for fact in document['facts']:
  for flow in fact.get('connections',[]):
   targets={e['entity_id']:e for e in named.get(norm(flow['target_name']),[])}
   if len(targets)!=1:notes.append(f'{fact["name"]}: connection to {flow["target_name"]} retained as unresolved because its entity reference is missing or ambiguous.');continue
   a=fact_entities[fact['record_id']];b=next(iter(targets.values()));rid=resource(flow['resource'],provenance=flow['source_ids'])
   if flow['direction']=='input':a,b=b,a
   key=(a['entity_id'],b['entity_id'],rid);match=next((l for l in links if (l['source_entity_id'],l['target_entity_id'],l['resource_id'])==key),None)
   if match is None:link(a,b,rid);match=links[-1]
   match.update({k:flow[k] for k in ['flow_rate','probability','delay','gate_condition']});match['provenance']={'source_ids':flow['source_ids'],'mapping_confidence':.7}
   relations.append({'type':'connection','name':f'{a["name"]} → {b["name"]}','resource_id':rid,'source_ids':flow['source_ids'],'note':flow['gate_condition'] or 'Explicit directed resource flow.'})
 for (rid,crid),pair in roundtrips.items():
  if 'buy' not in pair or 'sell' not in pair:continue
  buy,bent=pair['buy'];sell,sent=pair['sell'];loss=buy['value']-sell['value']
  relations.append({'type':'round_trip','name':f'Buy then sell {resources[rid]["name"]}','resource_id':crid,'net_loss':loss,'steps':[bent['entity_id'],sent['entity_id']],'source_ids':[buy['source_id'],sell['source_id']],'note':'Derived net change for one buy/sell cycle; not an additional sink or observed transaction frequency.'})
 for rid,items in weights.items():
  for weight in items:
   for bag,cap in bags:
    if weight['inventory'] and norm(weight['inventory'])!=norm(bag['name']):continue
    if cap['value'] and cap['unit']==weight['unit']:
     relations.append({'type':'capacity_use','name':f'{resources[rid]["name"]} in {bag["name"]}','resource_id':rid,'inventory_entity_id':bag['entity_id'],'weight':weight['value'],'capacity':cap['value'],'fraction':weight['value']/cap['value'],'condition':'If this inventory is equipped and stores one item; occupancy is not observed.','compatibility':'Explicit container reference' if weight['inventory'] else 'Possible pairing inferred from matching weight units; compatibility is unconfirmed.','source_ids':[weight['source_id'],cap['source_id']]})
 # Multiple canonical files are namespaced without inventing joins between games.
 for item in parsed.get('canonical_files',[]):
  eco=copy.deepcopy(item['economy']);prefix=f'canonical_{item["file_index"]}_';rs={r['resource_id']:prefix+r['resource_id'] for r in eco['resources']};es={e['entity_id']:prefix+e['entity_id'] for e in eco['entities']}
  for r in eco['resources']:r['resource_id']=rs[r['resource_id']];resources[r['resource_id']]=r
  for e in eco['entities']:
   old=e['entity_id'];e['entity_id']=es[old];e['affected_resources']=[rs[r] for r in e.get('affected_resources',[])];entities.append(e)
   for k,v in e.get('parameters',{}).items():
    if isinstance(v,(int,float)):bindings.append({'file_index':item['file_index'],'entity_id':e['entity_id'],'original_entity_id':old,'parameter':k,'original':v,'scale':1,'transform':'canonical','source_id':prefix+old+'.'+k})
  for l in eco.get('links',[]):l['link_id']=prefix+l['link_id'];l['source_entity_id']=es[l['source_entity_id']];l['target_entity_id']=es[l['target_entity_id']];l['resource_id']=rs[l['resource_id']];links.append(l)
  for key,value in eco.get('constraints',{}).items():
   if key not in constraints:constraints[key]=value
   elif constraints[key]!=value:constraint_conflicts.append({'file_index':item['file_index'],'key':key,'value':value})
 if not entities:raise ValueError('No resources or entities could be identified. The source tables are retained for an LLM-enabled import.')
 notes=list(dict.fromkeys(notes))
 notes.append('Automatic relationships are design interpretations. Unknown rates, currency identities and recipe quantities remain unknown.')
 eco={'schema_id':'px_economy_v2','game_id':project_id,'economy_id':project_id+'_v2','game_title':title,'archetype_set':'deep_progression','resources':list(resources.values()),'entities':entities,'links':links,'constraints':{},'data_quality':{'mapping_confidence':.7},'metadata':{'automatic_import':True,'interpretation_method':document['method'],'provisional':False,'relationships':relations,'notes':notes,'canonical_sources':[{'file_index':c['file_index'],'metadata':c['economy'].get('metadata',{})} for c in parsed.get('canonical_files',[])],'source_attributes':document['facts']}}
 eco['constraints']=constraints;eco['metadata']['constraint_conflicts']=constraint_conflicts
 # Conflicting files can still be assessed; never silently relax a hard constraint.
 if constraint_conflicts:eco['metadata']['balancing_block']='Conflicting constraints were retained from the canonical files; balancing needs a consistent constraint set.'
 # One source cell must not be independently tuned through multiple aliases.
 source_counts={}
 for b in bindings:source_counts[b['source_id']]=source_counts.get(b['source_id'],0)+1
 ents={e['entity_id']:e for e in entities}
 for b in bindings:
  if source_counts[b['source_id']]>1:
   e=ents[b['entity_id']];k=b['parameter'];v=e['parameters'][k];e['constraints'][k+'_min']=v;e['constraints'][k+'_max']=v
   notes.append('A repeated source cell is represented more than once; its numeric value is held fixed during automatic search.')
 if any(c['economy'].get('metadata',{}).get('provisional') for c in parsed.get('canonical_files',[])):eco['metadata']['provisional']=True
 # Preserve single-canonical identity, semantics and hard constraints exactly.
 if parsed.get('canonical'):
  eco=copy.deepcopy(parsed['canonical']['economy']);eco.setdefault('metadata',{})['automatic_import']=True;eco['metadata']['interpretation_method']='canonical';bindings=[]
 return eco,bindings,notes

def refresh_relationships(economy,bindings):
 """Recompute derived views from current parameters after accepted changes."""
 ents={e['entity_id']:e for e in economy['entities']};by_source={b['source_id']:ents[b['entity_id']]['parameters'].get(b['parameter']) for b in bindings if b.get('source_id') and b['entity_id'] in ents}
 rels=economy.get('metadata',{}).get('relationships',[])
 for r in rels:
  if r['type']=='exchange':
   e=ents[r['entity_id']];sale='selling' in e.get('functional_roles',[])
   r['outputs' if sale else 'inputs'][0]['quantity']=e['parameters'].get('cost_amount')
  elif r['type']=='round_trip':
   a,b=(ents[x] for x in r['steps']);r['net_loss']=a['parameters']['cost_amount']-b['parameters']['cost_amount']
  elif r['type']=='capacity_use':
   cap=ents[r['inventory_entity_id']]['parameters'].get('capacity');r['capacity']=cap;r['fraction']=r['weight']/cap if cap else None
 return economy
